import io, os, re, json, base64, datetime
import streamlit as st
from openai import OpenAI
from PIL import Image as PILImage
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as WA
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT, TA_CENTER
from templates import TEMPLATES, JENIS, brief_from_fields
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable

st.set_page_config(page_title="Ampera Scribe", page_icon="💬", layout="centered")

# ---------- Konfigurasi ----------
# Provider gratis (format OpenAI-compatible). Dicoba berurutan; key disimpan di Streamlit Secrets.
PROVIDERS = [
    ("Groq", "https://api.groq.com/openai/v1", "llama-3.3-70b-versatile", "GROQ_API_KEY"),
    ("Plugsky", "https://plugsky.com/v1", "plugsky-lite", "PLUGSKY_API_KEY"),  # bisa ditimpa lewat Secrets
]
PAPER = {"A4": (21, 29.7), "Letter": (21.59, 27.94), "Legal": (21.59, 35.56)}
FONTS = {"Times New Roman": ("Times-Roman", "Times-Bold"), "Arial": ("Helvetica", "Helvetica-Bold"),
         "Calibri": ("Helvetica", "Helvetica-Bold")}
PAGE_FMT = {"Tidak ada": "", "1": "{p}", "Halaman 1": "Halaman {p}",
            "Halaman 1 dari N": "Halaman {p} dari {n}", "1 / N": "{p} / {n}"}
ALIGN = {"Rata kiri-kanan": (TA_JUSTIFY, WA.JUSTIFY), "Rata kiri": (TA_LEFT, WA.LEFT)}
POS = ["Kiri", "Tengah", "Kanan"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus",
         "September", "Oktober", "November", "Desember"]
SIGN_W = 7  # lebar blok tanda tangan (cm)
APP_VER = "v5 - tata letak rapi"


def secret(k):
    try:
        return st.secrets[k]
    except Exception:
        return os.environ.get(k, "")


def tgl(d):
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def esc(t):
    return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


KV = re.compile(r"^([A-Za-z][\w./() -]{1,34}?)\s*:\s+(\S.*)$")


def blocks(text):
    """Pecah teks jadi elemen: ('h',judul) ('li',butir) ('kv',[(label,nilai)]) ('p',paragraf)."""
    out = []
    for b in re.split(r"\n\s*\n", (text or "").strip()):
        if not b.strip():
            continue
        par, kv = [], []

        def flush_par():
            if par:
                out.append(("p", " ".join(par)))
                par.clear()

        def flush_kv():
            if kv:
                out.append(("kv", list(kv)))
                kv.clear()

        for ln in b.splitlines():
            ln = ln.strip()
            if not ln:
                continue
            if ln.startswith("# "):
                flush_par(); flush_kv()
                out.append(("h", ln[2:].strip()))
            elif ln.startswith("- "):
                flush_par(); flush_kv()
                out.append(("li", ln[2:].strip()))
            elif KV.match(ln) and len(ln) < 90:
                flush_par()
                m = KV.match(ln)
                kv.append((m.group(1).strip(), m.group(2).strip()))
            else:
                flush_kv()
                par.append(ln)
        flush_par(); flush_kv()
    return out

# ---------- AI ----------
KEYS = ("judul", "pembuka", "isi", "penutup")


def _json_blob(txt):
    """Ambil objek JSON pertama yang kurung kurawalnya seimbang (abaikan ```fence```)."""
    txt = re.sub(r"```[a-zA-Z]*", "", txt or "").replace("```", "")
    start = txt.find("{")
    while start != -1:
        depth, in_str, esc_ch = 0, False, False
        for i in range(start, len(txt)):
            c = txt[i]
            if in_str:
                if esc_ch:
                    esc_ch = False
                elif c == "\\":
                    esc_ch = True
                elif c == '"':
                    in_str = False
            elif c == '"':
                in_str = True
            elif c == "{":
                depth += 1
            elif c == "}":
                depth -= 1
                if depth == 0:
                    return txt[start:i + 1]
        start = txt.find("{", start + 1)
    return None


REASON_PAT = re.compile(
    r"\b(we need to|we should|the user|let's craft|let us|i should|the spec|"
    r"as an ai|output json|valid json|must follow template|acceptable\.|probably okay)\b", re.I)


def _strip_reasoning(txt):
    """Buang blok berpikir model (<think>...</think>, 'Analysis:' dsb) dan pagar kode."""
    t = txt or ""
    t = re.sub(r"<(think|thinking|reasoning|scratchpad)>.*?</\1>", "", t, flags=re.S | re.I)
    t = re.sub(r"<(think|thinking|reasoning|scratchpad)>.*", "", t, flags=re.S | re.I)
    t = re.sub(r"```[a-zA-Z]*", "", t).replace("```", "")
    return t.strip()


def _salvage_json(txt):
    """JSON terpotong (kehabisan token): tutup string/kurung yang menggantung."""
    start = txt.find('{"')
    if start == -1:
        start = txt.find("{")
    if start == -1:
        return None
    t = txt[start:]
    depth, in_str, esc_ch = 0, False, False
    for c in t:
        if in_str:
            if esc_ch:
                esc_ch = False
            elif c == "\\":
                esc_ch = True
            elif c == '"':
                in_str = False
        elif c == '"':
            in_str = True
        elif c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
    if depth <= 0 and not in_str:
        return None
    t = t.rstrip().rstrip(",")
    if in_str:
        t += '"'
    return t + "}" * max(depth, 0)


def _from_text(txt):
    """Fallback: model balas teks biasa -> pecah jadi judul/pembuka/isi/penutup.
    Menolak teks yang jelas-jelas catatan berpikir model, bukan isi dokumen."""
    t = (txt or "").strip()
    if not t:
        return None
    if REASON_PAT.search(t[:1500]) or t.lower().count("json") >= 2:
        raise ValueError("model membalas catatan berpikir, bukan dokumen")
    par = [p.strip() for p in re.split(r"\n\s*\n", t) if p.strip()]
    judul = ""
    if par and len(par[0]) < 120 and "\n" not in par[0]:
        judul = par.pop(0).lstrip("# ").strip()
    pembuka = par.pop(0) if par else ""
    penutup = par.pop() if len(par) > 1 else ""
    return {"judul": judul, "pembuka": pembuka, "isi": "\n\n".join(par), "penutup": penutup}


def _normalize(data):
    out = {}
    for k in KEYS:
        v = data.get(k, "")
        if isinstance(v, list):
            v = "\n\n".join(map(str, v))
        elif isinstance(v, dict):
            v = "\n\n".join(f"# {a}\n{b}" for a, b in v.items())
        out[k] = str(v or "")
    return out


PLACEHOLDER = re.compile(r"^[\[\(<]{1,2}[^\]\)>]{0,60}[\]\)>]{1,2}[.,]?$")


def _esc_ctrl(t):
    """Escape newline/tab mentah yang ada DI DALAM string JSON (bikin json.loads gagal)."""
    out, in_str, esc_ch = [], False, False
    for c in t:
        if in_str:
            if esc_ch:
                esc_ch = False
            elif c == "\\":
                esc_ch = True
            elif c == '"':
                in_str = False
            elif c == "\n":
                out.append("\\n"); continue
            elif c == "\r":
                continue
            elif c == "\t":
                out.append("\\t"); continue
        elif c == '"':
            in_str = True
        out.append(c)
    return "".join(out)


def _loads(t):
    for cand in (t, _esc_ctrl(t)):
        try:
            return json.loads(cand)
        except Exception:
            pass
    return None


def _deep_find(obj, depth=0):
    """Cari dict {judul,pembuka,isi,penutup} walau terbungkus/di-encode berlapis."""
    if depth > 4:
        return None
    if isinstance(obj, str):
        t = obj.strip()
        if not (t.startswith("{") or t.startswith('"') or '\\"judul' in t or '"judul' in t):
            return None
        for cand in (t, _json_blob(t), t.replace('\\"', '"').replace("\\\\n", "\\n")):
            if not cand:
                continue
            v = _loads(cand)
            if v is not None and not isinstance(v, str):
                r = _deep_find(v, depth + 1)
                if r:
                    return r
            elif isinstance(v, str) and v != t:
                r = _deep_find(v, depth + 1)
                if r:
                    return r
            blob = _json_blob(cand)
            if blob and blob != cand:
                v2 = _loads(blob)
                if v2 is not None:
                    r = _deep_find(v2, depth + 1)
                    if r:
                        return r
        return None
    if isinstance(obj, list):
        for x in obj:
            r = _deep_find(x, depth + 1)
            if r:
                return r
        return None
    if isinstance(obj, dict):
        if any(k in obj for k in KEYS):
            return obj
        # kasus {"{\"judul\":...}": ""} -> JSON tersembunyi di key
        for k, v in obj.items():
            for cand in (k, v):
                r = _deep_find(cand, depth + 1)
                if r:
                    return r
    return None


def _tidy(v):
    """Rapikan satu nilai teks dari AI."""
    t = str(v or "")
    if "\\n" in t and "\n" not in t:          # newline masih berupa literal \n
        t = t.replace("\\n", "\n")
    t = t.replace("\\t", " ").replace('\\"', '"').replace("\\'", "'")
    t = t.strip()
    if t.startswith("{") and t.endswith("}"):
        t = t[1:-1].strip()
    if len(t) > 1 and t[0] == t[-1] and t[0] in "\"'":
        t = t[1:-1].strip()
    t = re.sub(r"^(judul|pembuka|isi|penutup)\s*:\s*", "", t, flags=re.I)
    out = []
    for ln in t.splitlines():
        s_ = ln.strip()
        if not s_:
            out.append("")
            continue
        if PLACEHOLDER.match(s_):            # buang [Penanggung Jawab], [Tanggal], dst.
            continue
        s_ = re.sub(r"^#{2,}\s*", "# ", s_)   # ## Judul -> # Judul
        if not s_.startswith("# "):
            s_ = re.sub(r"^[*+•]\s+", "- ", s_)
            s_ = re.sub(r"^\d+[.)]\s+", "- ", s_)
        out.append(s_)
    t = "\n".join(out)
    t = re.sub(r"\n{3,}", "\n\n", t)
    t = re.sub(r"[ \t]{2,}", " ", t)
    return t.strip()


def _split_sections(data):
    """Pindahkan sub-judul '# Pembuka' / '# Penutup' di dalam isi ke field-nya sendiri."""
    isi, buang = [], None
    blok = re.split(r"\n(?=# )", data["isi"])
    sisa = []
    for b in blok:
        head = b.split("\n", 1)[0].lstrip("# ").strip().lower()
        body = b.split("\n", 1)[1].strip() if "\n" in b else ""
        if head in ("pembuka", "salam pembuka", "pendahuluan surat") and not data["pembuka"]:
            data["pembuka"] = body
        elif head in ("penutup", "salam penutup", "kata penutup"):
            if not data["penutup"]:
                data["penutup"] = body
        else:
            sisa.append(b)
    data["isi"] = "\n\n".join(x.strip() for x in sisa if x.strip())
    return data


def _parse(raw):
    txt = _strip_reasoning(raw)
    for cand in (txt, _json_blob(txt), _salvage_json(txt)):
        if not cand:
            continue
        data = _deep_find(cand) or _deep_find(_loads(cand) if _loads(cand) is not None else None)
        if data:
            data = {k: _tidy(v) for k, v in _normalize(data).items()}
            if (data["isi"] + data["pembuka"]).strip():
                return _split_sections(data)
    data = _from_text(txt)
    if data:
        data = {k: _tidy(v) for k, v in _normalize(data).items()}
        return _split_sections(data)
    raise ValueError("balasan AI kosong / tidak bisa dibaca")


def generate(doc_type, brief, tone, length, tpl_nama="", outline=None):
    sys_msg = (
        "Kamu penulis dokumen resmi berbahasa Indonesia. "
        "JANGAN menulis proses berpikir, analisis, komentar, atau bahasa Inggris. "
        "Balas LANGSUNG satu objek JSON valid, tanpa teks lain dan tanpa blok kode, "
        'dengan tepat 4 kunci: {"judul":"","pembuka":"","isi":"","penutup":""}. '
        "Keempat nilai wajib berupa string berbahasa Indonesia yang sudah siap cetak.\n"
        'Aturan "isi": awali tiap sub-judul dengan "# ", tiap butir daftar dengan "- ", '
        "pisahkan paragraf dengan satu baris kosong. "
        'Jangan mengulang isi "pembuka" atau "penutup" di dalam "isi". '
        "Jangan membuat sub-judul bernama Pembuka atau Penutup. "
        "Jangan mengarang nama, nomor, atau tanggal yang tidak diberikan; "
        "lewati saja bagian yang datanya tidak ada. "
        "Jangan menulis salam penutup bertanda tangan atau nama penanda tangan."
    )
    ker = ""
    if outline:
        ker = ("\nKerangka bagian untuk \"isi\" (urut, tulis sebagai sub-judul '# ', "
               "lewati yang datanya tidak ada):\n"
               + "\n".join(f"{i}. {s}" for i, s in enumerate(outline, 1)) + "\n")
    usr = (f"Jenis dokumen: {doc_type}\nTemplate: {tpl_nama or doc_type}\n"
           f"Gaya bahasa: {tone}\nPanjang: {length}\n{ker}\nBahan:\n{brief}\n\n"
           "Keluarkan HANYA JSON-nya sekarang.")
    msgs = [{"role": "system", "content": sys_msg}, {"role": "user", "content": usr}]
    fix = {"role": "user", "content": ("Balasan sebelumnya tidak terpakai. Jangan tulis penjelasan "
                                       "atau proses berpikir. Keluarkan HANYA objek JSON dengan kunci "
                                       "judul, pembuka, isi, penutup dalam bahasa Indonesia.")}

    errors, skipped = [], []
    for name, url, model, key in PROVIDERS:
        k = secret(key)
        url = secret(key.replace("API_KEY", "BASE_URL")) or url
        model = secret(key.replace("API_KEY", "MODEL")) or model
        if not (k and url and model):
            skipped.append(name)
            continue
        client = OpenAI(base_url=url, api_key=k, timeout=180, max_retries=1)
        # 3 percobaan: JSON mode -> polos -> polos + teguran
        attempts = (({"response_format": {"type": "json_object"}}, msgs),
                    ({}, msgs),
                    ({}, msgs + [fix]))
        for kw, mm in attempts:
            try:
                r = client.chat.completions.create(model=model, temperature=0.3,
                                                   max_tokens=4096, messages=mm, **kw)
                raw = (r.choices[0].message.content or "").strip()
            except Exception as e:
                errors.append(f"{name} ({model}) - {type(e).__name__}: {e}")
                continue
            try:
                return _parse(raw), name
            except Exception as e:
                errors.append(f"{name} ({model}) - {e}. Cuplikan: {raw[:100]!r}")

    if not errors:
        raise RuntimeError("Belum ada API key aktif. Isi GROQ_API_KEY di Streamlit Secrets "
                           f"(provider yang dilewati: {', '.join(skipped) or '-'}).")
    raise RuntimeError("Semua provider gagal:\n- " + "\n- ".join(dict.fromkeys(errors)))


# ---------- Word ----------
def _field(run, code):
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = code
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "end")
    for x in (a, i, b):
        run._r.append(x)


def _hr(par, double=False):
    pPr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); ln = OxmlElement("w:bottom")
    for k, v in (("val", "double" if double else "single"), ("sz", "12"),
                 ("space", "1"), ("color", "000000")):
        ln.set(qn("w:" + k), v)
    bd.append(ln); pPr.append(bd)


def build_docx(d):
    c, k, sg = d["cfg"], d["kop"], d["sign"]
    doc = Document(); s = doc.sections[0]
    pw, ph = PAPER[c["paper"]]
    s.page_width, s.page_height = Cm(pw), Cm(ph)
    s.left_margin = s.right_margin = Cm(c["mx"])
    s.top_margin = s.bottom_margin = Cm(c["my"])
    W = pw - 2 * c["mx"]
    n = doc.styles["Normal"]
    n.font.name = c["font"]; n.font.size = Pt(c["size"])
    n.element.rPr.rFonts.set(qn("w:eastAsia"), c["font"])
    n.paragraph_format.line_spacing = c["spacing"]
    n.paragraph_format.space_after = Pt(6)

    def P(text="", align=None, bold=False, size=None, after=None, container=None, underline=False):
        p = (container or doc).add_paragraph()
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text); r.bold = bold; r.underline = underline
            if size:
                r.font.size = Pt(size)
        return p

    if k["name"] or k["logo"]:
        t = doc.add_table(rows=1, cols=2); t.autofit = False
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Cm(3), Cm(W - 3)
        if k["logo"]:
            c0.paragraphs[0].add_run().add_picture(io.BytesIO(k["logo"]), width=Cm(2.3))
        first = c1.paragraphs[0]; first.alignment = WA.CENTER
        first.paragraph_format.space_after = Pt(0)
        r = first.add_run(k["name"].upper()); r.bold = True; r.font.size = Pt(c["size"] + 4)
        for line in (k["addr"], k["contact"]):
            if line:
                P(line, WA.CENTER, size=c["size"] - 1, after=0, container=c1)
        _hr(P(after=10), double=True)

    body_al = ALIGN[c["align"]][1]
    if d["meta"]:
        m = d["meta"]
        P(m["kota_tgl"], WA.RIGHT)
        for lab, val in (("Nomor", m["nomor"]), ("Lampiran", m["lampiran"]), ("Perihal", m["perihal"])):
            if val:
                p = P(f"{lab}\t: {val}", after=0)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(2.5))
        P()
        if m["tujuan"]:
            P("Yth. " + m["tujuan"], after=0); P("di tempat")
    else:
        p = P(d["judul"].upper(), WA.CENTER, bold=True, size=c["size"] + 2, underline=True)
        p.paragraph_format.space_after = Pt(14)

    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                p = P(txt, bold=True, after=4)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.keep_with_next = True
            elif kind == "li":
                p = doc.add_paragraph(txt, style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                p.alignment = body_al
            elif kind == "kv":
                wlab = min(max(len(a) for a, _ in txt) * 0.22 + 0.6, 6.5)
                t = doc.add_table(rows=0, cols=3); t.autofit = False
                for lab, val in txt:
                    row = t.add_row().cells
                    row[0].width, row[1].width, row[2].width = Cm(wlab), Cm(0.5), Cm(W - wlab - 0.5)
                    for cell, s_ in ((row[0], lab), (row[1], ":"), (row[2], val)):
                        pp = cell.paragraphs[0]
                        pp.paragraph_format.space_after = Pt(0)
                        pp.add_run(s_)
                P(after=4)
            else:
                P(txt, body_al)

    if sg["nama"]:
        left = {"Kiri": 0, "Tengah": (W - SIGN_W) / 2, "Kanan": W - SIGN_W}[sg["pos"]]
        right = {"Kiri": W - SIGN_W, "Tengah": (W - SIGN_W) / 2, "Kanan": 0}[sg["pos"]]

        def SP(*a, **kw):
            p = P(*a, **kw)
            p.paragraph_format.left_indent, p.paragraph_format.right_indent = Cm(left), Cm(right)
            return p

        P()
        SP(sg["kota_tgl"], WA.CENTER, after=0)
        SP(sg["jabatan"], WA.CENTER, after=0)
        if sg["img"]:
            p = SP(align=WA.CENTER, after=0)
            p.add_run().add_picture(io.BytesIO(sg["img"]), width=Cm(3.5))
        else:
            for _ in range(3):
                SP(after=0)
        SP(sg["nama"], WA.CENTER, bold=True, underline=True, after=0)
        if sg["nip"]:
            SP(sg["nip"], WA.CENTER, after=0)

    tpl = PAGE_FMT[c["pagefmt"]]
    if tpl:
        p = s.footer.paragraphs[0]
        p.alignment = {"Kiri": WA.LEFT, "Tengah": WA.CENTER, "Kanan": WA.RIGHT}[c["pagepos"]]
        for part in re.split(r"(\{p\}|\{n\})", tpl):
            if part == "{p}":
                _field(p.add_run(), "PAGE")
            elif part == "{n}":
                _field(p.add_run(), "NUMPAGES")
            elif part:
                p.add_run(part)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


# ---------- PDF ----------
def _img(data, w_cm):
    iw, ih = PILImage.open(io.BytesIO(data)).size
    return Image(io.BytesIO(data), width=w_cm * cm, height=w_cm * cm * ih / iw)


def _canvas(tpl, pos, font, size, pw, mx, my):
    class NC(rl_canvas.Canvas):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw); self._st = []

        def showPage(self):
            self._st.append(dict(self.__dict__)); self._startPage()

        def save(self):
            total = len(self._st)
            for x in self._st:
                self.__dict__.update(x)
                if tpl:
                    self.setFont(font, size - 1)
                    t, y = tpl.format(p=self._pageNumber, n=total), my / 2
                    if pos == "Kiri":
                        self.drawString(mx, y, t)
                    elif pos == "Kanan":
                        self.drawRightString(pw - mx, y, t)
                    else:
                        self.drawCentredString(pw / 2, y, t)
                super().showPage()
            super().save()
    return NC


def build_pdf(d):
    c, k, sg = d["cfg"], d["kop"], d["sign"]
    reg, bold = FONTS[c["font"]]
    pw, ph = [x * cm for x in PAPER[c["paper"]]]
    mx, my = c["mx"] * cm, c["my"] * cm
    W = pw - 2 * mx
    lead = c["size"] * c["spacing"] * 1.2
    S = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("fontName", reg), fontSize=kw.pop("fontSize", c["size"]),
                                          leading=kw.pop("leading", lead), spaceAfter=kw.pop("spaceAfter", 6), **kw)
    body = S("body", alignment=ALIGN[c["align"]][0])
    ctr = S("ctr", alignment=TA_CENTER, spaceAfter=0)
    story = []

    if k["name"] or k["logo"]:
        txt = [Paragraph(esc(k["name"].upper()), S("kn", fontName=bold, fontSize=c["size"] + 4, alignment=TA_CENTER,
                                            leading=(c["size"] + 4) * 1.2, spaceAfter=0))]
        for line in (k["addr"], k["contact"]):
            if line:
                txt.append(Paragraph(esc(line), S("kl", fontSize=c["size"] - 1, alignment=TA_CENTER, spaceAfter=0)))
        if k["logo"]:
            t = Table([[_img(k["logo"], 2.3), txt]], colWidths=[3 * cm, W - 3 * cm])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
            story.append(t)
        else:
            story += txt
        story += [Spacer(1, 4), HRFlowable(width="100%", thickness=1.6, color=colors.black),
                  Spacer(1, 1.5), HRFlowable(width="100%", thickness=0.6, color=colors.black),
                  Spacer(1, 12)]

    if d["meta"]:
        m = d["meta"]
        story.append(Paragraph(esc(m["kota_tgl"]), S("r", alignment=2)))
        rows = [[lab, ": " + val] for lab, val in (("Nomor", m["nomor"]), ("Lampiran", m["lampiran"]),
                                                    ("Perihal", m["perihal"])) if val]
        if rows:
            t = Table([[Paragraph(esc(a), body), Paragraph(esc(b), body)] for a, b in rows],
                      colWidths=[2.5 * cm, W - 2.5 * cm], hAlign="LEFT")
            t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 0)]))
            story += [t, Spacer(1, 10)]
        if m["tujuan"]:
            story += [Paragraph("Yth. " + esc(m["tujuan"]), S("y", spaceAfter=0)), Paragraph("di tempat", body)]
    else:
        story.append(Paragraph(f"<u>{esc(d['judul'].upper())}</u>",
                               S("t", fontName=bold, fontSize=c["size"] + 2, alignment=TA_CENTER,
                                 leading=(c["size"] + 2) * 1.3, spaceAfter=14)))

    head = S("hd", fontName=bold, spaceBefore=10, spaceAfter=4)
    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                story.append(Paragraph(esc(txt), head))
            elif kind == "li":
                story.append(Paragraph(esc(txt), S("li", alignment=ALIGN[c["align"]][0], spaceAfter=3,
                                                   leftIndent=14), bulletText="\u2022"))
            elif kind == "kv":
                wlab = min(max(len(a) for a, _ in txt) * 0.22 + 0.6, 6.5) * cm
                rows = [[Paragraph(esc(a), body), Paragraph(":", body), Paragraph(esc(b), body)]
                        for a, b in txt]
                t = Table(rows, colWidths=[wlab, 0.5 * cm, W - wlab - 0.5 * cm], hAlign="LEFT")
                t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                                       ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                       ("TOPPADDING", (0, 0), (-1, -1), 0),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story += [t, Spacer(1, 5)]
            else:
                story.append(Paragraph(esc(txt), body))

    if sg["nama"]:
        cell = [Spacer(1, 12), Paragraph(esc(sg["kota_tgl"]), ctr), Paragraph(esc(sg["jabatan"]), ctr)]
        cell.append(_img(sg["img"], 3.5) if sg["img"] else Spacer(1, 2 * cm))
        cell.append(Paragraph(f"<u><b>{esc(sg['nama'])}</b></u>", ctr))
        if sg["nip"]:
            cell.append(Paragraph(esc(sg["nip"]), ctr))
        t = Table([[cell]], colWidths=[SIGN_W * cm], hAlign={"Kiri": "LEFT", "Tengah": "CENTER", "Kanan": "RIGHT"}[sg["pos"]])
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        story.append(t)

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=(pw, ph), leftMargin=mx, rightMargin=mx, topMargin=my, bottomMargin=my).build(
        story, canvasmaker=_canvas(PAGE_FMT[c["pagefmt"]], c["pagepos"], reg, c["size"], pw, mx, my))
    return buf.getvalue()


# ---------- Pratinjau HTML ----------
def _b64(data):
    return "data:image/png;base64," + base64.b64encode(data).decode()


def build_html(d):
    """Render dokumen jadi HTML yang meniru tata letak halaman (untuk pratinjau)."""
    c, k, sg = d["cfg"], d["kop"], d["sign"]
    fam = {"Times New Roman": "'Times New Roman', Times, serif",
           "Arial": "Arial, Helvetica, sans-serif",
           "Calibri": "Calibri, Carlito, sans-serif"}[c["font"]]
    pw = PAPER[c["paper"]][0]
    just = "justify" if c["align"].startswith("Rata kiri-kanan") else "left"
    o = []

    if k["name"] or k["logo"]:
        logo = f'<img src="{_b64(k["logo"])}" style="width:2.3cm">' if k["logo"] else ""
        sub = "".join(f'<div class="kl">{esc(x)}</div>' for x in (k["addr"], k["contact"]) if x)
        o.append(f'<div class="kop"><div class="klogo">{logo}</div>'
                 f'<div class="ktxt"><div class="kn">{esc(k["name"].upper())}</div>{sub}</div></div><hr>')

    if d["meta"]:
        m = d["meta"]
        o.append(f'<div class="right">{esc(m["kota_tgl"])}</div>')
        rows = "".join(f'<tr><td class="lab">{esc(a)}</td><td>: {esc(b)}</td></tr>'
                       for a, b in (("Nomor", m["nomor"]), ("Lampiran", m["lampiran"]),
                                    ("Perihal", m["perihal"])) if b)
        if rows:
            o.append(f'<table class="meta">{rows}</table>')
        if m["tujuan"]:
            o.append(f'<div class="tuj">Yth. {esc(m["tujuan"])}<br>di tempat</div>')
    else:
        o.append(f'<div class="judul">{esc(d["judul"].upper())}</div>')

    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                o.append(f'<div class="h">{esc(txt)}</div>')
            elif kind == "li":
                o.append(f'<ul><li>{esc(txt)}</li></ul>')
            elif kind == "kv":
                rows = "".join(f'<tr><td class="kvl">{esc(a)}</td><td class="kvc">:</td>'
                               f'<td>{esc(b)}</td></tr>' for a, b in txt)
                o.append(f'<table class="kv">{rows}</table>')
            else:
                o.append(f'<p>{esc(txt)}</p>')

    if sg["nama"]:
        al = {"Kiri": "flex-start", "Tengah": "center", "Kanan": "flex-end"}[sg["pos"]]
        img = (f'<img src="{_b64(sg["img"])}" style="width:3.5cm">' if sg["img"]
               else '<div style="height:2cm"></div>')
        o.append(f'<div class="sgwrap" style="justify-content:{al}"><div class="sg">'
                 f'<div>{esc(sg["kota_tgl"])}</div><div>{esc(sg["jabatan"])}</div>{img}'
                 f'<div class="nm">{esc(sg["nama"])}</div>'
                 f'<div>{esc(sg["nip"])}</div></div></div>')

    return f"""<style>
.paper{{background:#fff;color:#111;font-family:{fam};font-size:{c["size"]}pt;
line-height:{c["spacing"]};padding:{c["my"]}cm {c["mx"]}cm;width:{pw}cm;max-width:100%;
box-sizing:border-box;margin:0 auto;box-shadow:0 2px 14px rgba(0,0,0,.25);border-radius:4px;}}
.paper p{{text-align:{just};margin:0 0 6pt;}}
.paper hr{{border:0;border-top:1.6px solid #000;border-bottom:.6px solid #000;height:2px;margin:5px 0 12px;}}
.paper .kop{{display:flex;align-items:center;gap:.4cm;}}
.paper .klogo{{width:3cm;text-align:center;}}
.paper .ktxt{{flex:1;text-align:center;}}
.paper .kn{{font-weight:700;font-size:{c["size"] + 4}pt;}}
.paper .kl{{font-size:{c["size"] - 1}pt;}}
.paper .right{{text-align:right;margin-bottom:6pt;}}
.paper .meta td{{padding:0;vertical-align:top;}}
.paper .meta .lab{{width:2.5cm;}}
.paper .tuj{{margin:10pt 0;}}
.paper .judul{{text-align:center;font-weight:700;text-decoration:underline;font-size:{c["size"] + 2}pt;margin-bottom:14pt;}}
.paper .h{{font-weight:700;margin:10pt 0 4pt;}}
.paper table.kv{{margin:0 0 5pt;border-collapse:collapse;}}
.paper table.kv td{{padding:0 0 1pt;vertical-align:top;}}
.paper table.kv .kvl{{white-space:nowrap;padding-right:.3cm;}}
.paper table.kv .kvc{{width:.4cm;}}
.paper ul{{margin:0 0 4pt 0;padding-left:1.2em;}}
.paper .sgwrap{{display:flex;margin-top:14pt;}}
.paper .sg{{width:7cm;text-align:center;}}
.paper .sg .nm{{font-weight:700;text-decoration:underline;}}
</style><div class="paper">{"".join(o)}</div>"""

# ---------- UI ----------
CSS = """
<style>
.block-container {padding-top: 2rem; max-width: 1000px;}
div[data-testid="stChatMessage"] {background: rgba(128,128,128,.07); border-radius: 14px; padding: .6rem .9rem;}
.go-btn button {width: 100%; height: 3.2rem; font-size: 1.15rem; font-weight: 700; border-radius: 14px;}
.hint {opacity:.65; font-size:.88rem;}
section[data-testid="stSidebar"] .stExpander {border-radius: 12px;}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

ss = st.session_state
ss.setdefault("chat", [])       # [{"role": "user"/"assistant", "text": str}]
ss.setdefault("ready", False)
for k in ("judul", "pembuka", "isi", "penutup"):
    ss.setdefault(k, "")


def reset_doc():
    ss.ready = False


# ================= SIDEBAR: semua pengaturan =================
with st.sidebar:
    st.markdown("## ⚙️ Pengaturan dokumen")
    st.caption(f"Ampera Scribe {APP_VER}")

    # --- Langkah 1: jenis file ---
    st.markdown("**Langkah 1 — Jenis file**")
    doc_type = st.selectbox("Jenis dokumen", JENIS, label_visibility="collapsed",
                            key="doc_type", on_change=reset_doc)

    # --- Langkah 2: template ---
    st.markdown("**Langkah 2 — Template**")
    opsi = TEMPLATES[doc_type]
    tpl_nama = st.selectbox("Template", [t["nama"] for t in opsi], label_visibility="collapsed",
                            key=f"tpl_{doc_type}", on_change=reset_doc)
    tpl = next(t for t in opsi if t["nama"] == tpl_nama)
    st.caption(f"📌 {tpl['desc']}")
    with st.popover("Lihat kerangka template", use_container_width=True):
        st.markdown("\n".join(f"{i}. {s}" for i, s in enumerate(tpl["outline"], 1)))

    st.divider()

    # --- Langkah 3: isian dokumen ---
    st.markdown("**Langkah 3 — Isi dokumen**")
    with st.expander("📝 Judul & identitas", expanded=True):
        judul_in = st.text_input("Judul dokumen", tpl["judul"], key=f"jd_{tpl['id']}")
        kota = st.text_input("Kota", "Bandar Lampung")
        tanggal = st.date_input("Tanggal", datetime.date.today())

    is_surat = tpl["surat"]
    nomor = lampiran = perihal = tujuan = ""
    if is_surat:
        with st.expander("✉️ Nomor & tujuan surat", expanded=True):
            nomor = st.text_input("Nomor surat")
            lampiran = st.text_input("Lampiran")
            perihal = st.text_input("Perihal (kosong = ikut judul)")
            tujuan = st.text_input("Ditujukan kepada")

    vals = {}
    with st.expander(f"🧩 Data khusus: {tpl['nama']}", expanded=True):
        for f in tpl["fields"]:
            wk = f"{tpl['id']}__{f['key']}"
            if f["type"] == "area":
                vals[f["key"]] = st.text_area(f["label"], placeholder=f["ph"], height=90, key=wk)
            elif f["type"] == "date":
                vals[f["key"]] = str(st.date_input(f["label"], datetime.date.today(), key=wk))
            else:
                vals[f["key"]] = st.text_input(f["label"], placeholder=f["ph"], key=wk)

    st.divider()

    # --- Langkah 4: tampilan dokumen ---
    st.markdown("**Langkah 4 — Tampilan & format**")
    with st.expander("🏛️ Kop surat"):
        kop_name = st.text_input("Nama instansi / organisasi")
        kop_addr = st.text_area("Alamat", height=70)
        kop_contact = st.text_input("Telp / email / website")
        logo = st.file_uploader("Logo (PNG/JPG)", type=["png", "jpg", "jpeg"])

    with st.expander("🎨 Gaya tulisan"):
        tone = st.selectbox("Gaya bahasa", ["Formal", "Semi-formal", "Ringkas dan padat"])
        length = st.selectbox("Target panjang", ["1 halaman", "2 halaman", "3 halaman atau lebih"])

    with st.expander("📐 Layout halaman"):
        paper = st.selectbox("Ukuran kertas", list(PAPER))
        mx = st.slider("Margin kiri-kanan (cm)", 1.0, 4.0, 2.5, 0.1)
        my = st.slider("Margin atas-bawah (cm)", 1.5, 4.0, 2.5, 0.1)
        font = st.selectbox("Font", list(FONTS))
        size = st.slider("Ukuran font", 9, 16, 12)
        spacing = st.slider("Spasi baris", 1.0, 2.0, 1.15, 0.05)
        align = st.selectbox("Perataan teks", list(ALIGN))
        pagefmt = st.selectbox("Nomor halaman", list(PAGE_FMT), index=3)
        pagepos = st.selectbox("Posisi nomor halaman", POS, index=1)

    with st.expander("✍️ Tanda tangan"):
        jabatan = st.text_input("Jabatan penandatangan")
        nama = st.text_input("Nama penandatangan")
        nip = st.text_input("NIP / info tambahan")
        sign_pos = st.selectbox("Posisi tanda tangan", POS, index=2)
        sign_img = st.file_uploader("Gambar TTD / stempel", type=["png", "jpg", "jpeg"])

    st.divider()
    if st.button("🧹 Mulai obrolan baru", use_container_width=True):
        ss.chat = []
        ss.ready = False
        for k in ("judul", "pembuka", "isi", "penutup"):
            ss[k] = ""
        st.rerun()

# ================= HALAMAN UTAMA: ruang chat =================
st.title("💬 Ampera Scribe")
st.caption("Pilih jenis file dan template di sidebar → isi datanya → tambahkan catatan bebas di chat → "
           "tekan tombol besar di bawah.")

isian = brief_from_fields(tpl, vals)
catatan = "\n".join(m["text"] for m in ss.chat if m["role"] == "user").strip()
brief = "\n".join(x for x in (f"Judul: {judul_in}" if judul_in else "", isian,
                              ("Catatan tambahan dari pengguna:\n" + catatan) if catatan else "") if x)

if not ss.chat:
    with st.chat_message("assistant", avatar="📄"):
        n_tpl = sum(len(v) for v in TEMPLATES.values())
        st.markdown(
            f"Halo! Template aktif: **{tpl['nama']}** ({doc_type}) — satu dari {n_tpl} template yang tersedia.\n\n"
            f"_{tpl['desc']}_\n\n"
            "Isi data di sidebar, dan kalau ada detail yang tidak ada kolomnya, tulis saja di chat bawah. "
            "Kalau sudah siap, klik **let's go!! Buat file nya**."
        )

for m in ss.chat:
    with st.chat_message(m["role"], avatar="🧑" if m["role"] == "user" else "📄"):
        st.markdown(m["text"])

if msg := st.chat_input("Tulis catatan tambahan atau permintaan revisi..."):
    ss.chat.append({"role": "user", "text": msg})
    ss.chat.append({"role": "assistant",
                    "text": "Oke, sudah kucatat ✅ Tambah detail lagi, atau langsung tekan **let's go!! Buat file nya**."})
    st.rerun()

with st.expander("👀 Ringkasan bahan yang akan dikirim ke AI"):
    st.code(brief or "(belum ada data)", language="text")

st.markdown('<div class="go-btn">', unsafe_allow_html=True)
go = st.button("🚀 let's go!! Buat file nya", type="primary", disabled=not isian and not catatan,
               use_container_width=True)
st.markdown("</div>", unsafe_allow_html=True)
if not isian and not catatan:
    st.markdown('<p class="hint">Isi minimal satu data di sidebar (atau tulis di chat) '
                'untuk mengaktifkan tombol.</p>', unsafe_allow_html=True)

if go:
    with st.spinner(f"AI sedang menyusun {tpl['nama']}..."):
        try:
            res, used = generate(doc_type, brief, tone, length, tpl["nama"], tpl["outline"])
            for k in ("judul", "pembuka", "isi", "penutup"):
                ss[k] = str(res.get(k, ""))
            if judul_in:
                ss["judul"] = judul_in
            ss.ready = True
            ss.chat.append({"role": "assistant",
                            "text": f"Selesai! **{ss['judul'] or tpl['nama']}** sudah jadi (via {used}). "
                                    "Cek dan edit hasilnya di bawah, lalu unduh Word atau PDF."})
            st.rerun()
        except Exception as e:
            st.error(f"Gagal: {e}")

# ================= Hasil =================
if ss.ready:
    st.divider()
    kt = f"{kota}, {tgl(tanggal)}"
    d = {
        "cfg": dict(paper=paper, mx=mx, my=my, font=font, size=size, spacing=spacing, align=align,
                    pagefmt=pagefmt, pagepos=pagepos),
        "kop": dict(name=kop_name, addr=kop_addr, contact=kop_contact, logo=logo.getvalue() if logo else None),
        "meta": dict(kota_tgl=kt, nomor=nomor, lampiran=lampiran, perihal=perihal or ss["judul"],
                     tujuan=tujuan) if is_surat else None,
        "sign": dict(jabatan=jabatan, nama=nama, nip=nip, pos=sign_pos, kota_tgl=kt,
                     img=sign_img.getvalue() if sign_img else None),
        "judul": ss["judul"], "pembuka": ss["pembuka"], "isi": ss["isi"], "penutup": ss["penutup"],
    }

    st.subheader("📄 Dokumen kamu")
    tab_prev, tab_edit = st.tabs(["👁️ Pratinjau", "✏️ Edit teks"])

    with tab_edit:
        st.caption("Setiap perubahan langsung terlihat di tab Pratinjau dan ikut terbawa saat diunduh.")
        st.text_input("Judul", key="judul")
        st.text_area("Pembuka", key="pembuka", height=110)
        st.text_area("Isi  —  '# ' sub-judul, '- ' butir, baris kosong = paragraf baru",
                     key="isi", height=380)
        st.text_area("Penutup", key="penutup", height=90)

    with tab_prev:
        st.caption("Tampilan mendekati hasil cetak. Perbaiki teksnya di tab Edit teks bila ada yang keliru.")
        st.markdown(build_html(d), unsafe_allow_html=True)

    st.write("")
    fn = re.sub(r"[^A-Za-z0-9]+", "-", (ss["judul"] or tpl["nama"]).strip()).strip("-").lower() or "dokumen"
    try:
        d1, d2 = st.columns(2)
        d1.download_button("⬇️ Unduh Word (.docx)", build_docx(d), f"{fn}.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True)
        d2.download_button("⬇️ Unduh PDF", build_pdf(d), f"{fn}.pdf", "application/pdf",
                           use_container_width=True)
    except Exception as e:
        st.error(f"Gagal membuat file: {e}")
