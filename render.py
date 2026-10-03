"""Renderer dokumen Ampera Scribe: Word (.docx), PDF, dan pratinjau HTML.

Struktur data dokumen (dict `d`):
    cfg    : paper, mx, my, font, size, spacing, align, pagefmt, pagepos
    kop    : name, addr, contact, logo  (None kalau tanpa kop)
    meta   : None | {kota_tgl, nomor, lampiran, hal, sifat, tujuan, up}
    memo   : None | {tanggal, dari, kepada, tembusan, subyek}
    judul  : judul dokumen ("" kalau pakai blok meta)
    nomor_judul : nomor surat yang dicetak di bawah judul
    pembuka / isi / penutup : teks (mini-markup)
    ttds   : [{label, nama, jabatan, nip, img}]
    ttd_pos, kota_tgl, tutup_tgl (kota|ditetapkan)
    verifikasi, tembusan[], nb, registrasi, materai, saksi[], slip

Kop surat: teks kop SELALU rata tengah pada seluruh lebar halaman. Logo (bila ada)
diletakkan di sisi kiri tanpa menggeser teks kop.
"""
import io, os, re, base64
from PIL import Image as PILImage
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as WA
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT, TA_CENTER
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image, Table,
                                TableStyle, HRFlowable, KeepTogether)

PAPER = {"A4": (21, 29.7), "Letter": (21.59, 27.94), "Legal": (21.59, 35.56), "F4": (21.5, 33)}
FONTS = {"Times New Roman": ("Times-Roman", "Times-Bold"),
         "Arial": ("Helvetica", "Helvetica-Bold"),
         "Calibri": ("Helvetica", "Helvetica-Bold")}
PAGE_FMT = {"Tidak ada": "", "1": "{p}", "Halaman 1": "Halaman {p}",
            "Halaman 1 dari N": "Halaman {p} dari {n}", "1 / N": "{p} / {n}"}
ALIGN = {"Rata kiri-kanan": (TA_JUSTIFY, WA.JUSTIFY), "Rata kiri": (TA_LEFT, WA.LEFT)}
POS = ["Kiri", "Tengah", "Kanan"]
SIGN_W = 7.0
MATERAI = "Materai\nRp 10.000"
KOP_LOGO_W = 2.6   # cm: lebar lajur logo di kiri DAN lajur kosong di kanan (supaya teks tetap di tengah)


# ---------------------------------------------------------------- teks Arab
AR_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]+"
                   r"(?:[\s\u0600-\u06FF\u064B-\u065F.,:;!?()-]*"
                   r"[\u0600-\u06FF\uFE70-\uFEFF])?")
AR_FONT = None
for _p in ("fonts/Amiri-Regular.ttf", "fonts/NotoNaskhArabic-Regular.ttf",
           "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
    if os.path.exists(_p):
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            pdfmetrics.registerFont(TTFont("ArabFont", _p))
            AR_FONT = ("ArabFont", os.path.basename(_p).rsplit("-", 1)[0]
                       .replace("Regular", "").strip() or "Amiri")
            break
        except Exception:
            pass
AR_DOCX_FONT = "Amiri" if os.path.exists("fonts/Amiri-Regular.ttf") else "DejaVu Sans"


def has_ar(t):
    return bool(AR_RE.search(t or ""))


def shape_ar(t):
    """Bentuk ulang + urutkan kanan-ke-kiri penggalan Arab agar terbaca benar."""
    if not has_ar(t):
        return t
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
    except Exception:
        return t
    return AR_RE.sub(lambda m: get_display(arabic_reshaper.reshape(m.group(0))), t)


def ar_markup(t):
    """Untuk PDF: bungkus penggalan Arab dengan font yang mendukung Arab."""
    if not (AR_FONT and has_ar(t)):
        return t
    return AR_RE.sub(lambda m: f'<font name="{AR_FONT[0]}">{m.group(0)}</font>', t)


def esc(t):
    t = (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return ar_markup(shape_ar(t))


# ---------------------------------------------------------------- parser isi
KV = re.compile(r"^([A-Za-z][\w./() -]{1,34}?)\s*:\s+(\S.*)$")
NO = re.compile(r"^(\d{1,2}|[a-z])[.)]\s+(\S.*)$")


def blocks(text):
    """Pecah teks jadi elemen:
    ('h',judul) ('li',butir) ('no',(label,teks)) ('kv',[(a,b)]) ('tbl',[[sel]]) ('p',par)
    """
    out = []
    for b in re.split(r"\n\s*\n", (text or "").strip()):
        if not b.strip():
            continue
        par, kv, tbl = [], [], []

        def flush_p():
            if par:
                out.append(("p", " ".join(par))); par.clear()

        def flush_kv():
            if kv:
                out.append(("kv", list(kv))); kv.clear()

        def flush_tbl():
            if tbl:
                out.append(("tbl", list(tbl))); tbl.clear()

        for ln in b.splitlines():
            ln = ln.strip()
            if not ln:
                continue
            if ln.startswith("|") or (ln.count("|") >= 1 and not ln.startswith("# ")
                                      and len(ln.split("|")) > 1 and not KV.match(ln)):
                flush_p(); flush_kv()
                cells = [c.strip() for c in ln.strip("|").split("|")]
                if any(cells):
                    tbl.append(cells)
                continue
            flush_tbl()
            if ln.startswith("# "):
                flush_p(); flush_kv()
                out.append(("h", ln[2:].strip()))
            elif ln.startswith("- "):
                flush_p(); flush_kv()
                out.append(("li", ln[2:].strip()))
            elif NO.match(ln):
                flush_p(); flush_kv()
                m = NO.match(ln)
                out.append(("no", (m.group(1) + ".", m.group(2).strip())))
            elif KV.match(ln) and len(ln) < 95:
                flush_p()
                m = KV.match(ln)
                kv.append((m.group(1).strip(), m.group(2).strip()))
            else:
                flush_kv()
                par.append(ln)
        flush_p(); flush_kv(); flush_tbl()
    return out


def parse_list(txt):
    return [l.strip() for l in (txt or "").splitlines() if l.strip()]


def parse_table(txt):
    rows = []
    for ln in parse_list(txt):
        rows.append([c.strip() for c in ln.strip("|").split("|")])
    return rows


def tujuan_lines(m):
    """Baris alamat tujuan tanpa 'di Tempat' di ujungnya (renderer menambahkannya sendiri,
    jadi tidak dobel kalau pengguna sudah mengetiknya)."""
    lines = parse_list(m.get("tujuan"))
    while lines and re.sub(r"[\s.,:;]+", " ", lines[-1]).strip().lower() == "di tempat":
        lines.pop()
    return lines


def _heads(row, ncol, kota_line, total, materai):
    """Baris teks di atas ruang tanda tangan, disamakan tingginya antar kolom."""
    heads = []
    for idx in range(ncol):
        sg = row[idx] if idx < len(row) else None
        lines = []
        if sg:
            if kota_line and (idx == len(row) - 1 or total == 1):
                lines.append(kota_line)
            for key in ("label", "jabatan"):
                if sg.get(key):
                    lines.append(sg[key])
            if materai and idx == len(row) - 1:
                lines.append(MATERAI.replace("\n", " "))
        heads.append(lines)
    tinggi = max((len(x) for x in heads), default=0)
    return [[""] * (tinggi - len(x)) + x for x in heads]


# ==================================================================== WORD
def _field(run, code):
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = code
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "end")
    for x in (a, i, b):
        run._r.append(x)


def _hr(par, double=False):
    pPr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); ln = OxmlElement("w:bottom")
    for k, v in (("val", "double" if double else "single"), ("sz", "12"),
                 ("space", "1"), ("color", "000000")):
        ln.set(qn("w:" + k), v)
    bd.append(ln); pPr.append(bd)


def _borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000")
        borders.append(e)
    pr.append(borders)


def _valign_center(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign"); v.set(qn("w:val"), "center")
    tcPr.append(v)


def build_docx(d):
    c, k = d["cfg"], d["kop"]
    doc = Document(); s = doc.sections[0]
    pw, ph = PAPER[c["paper"]]
    s.page_width, s.page_height = Cm(pw), Cm(ph)
    s.left_margin = s.right_margin = Cm(c["mx"])
    s.top_margin = s.bottom_margin = Cm(c["my"])
    W = pw - 2 * c["mx"]
    n = doc.styles["Normal"]
    n.font.name = c["font"]; n.font.size = Pt(c["size"])
    n.element.rPr.rFonts.set(qn("w:eastAsia"), c["font"])
    n.paragraph_format.line_spacing = c["spacing"]
    n.paragraph_format.space_after = Pt(6)
    body_al = ALIGN[c["align"]][1]

    def P(text="", align=None, bold=False, size=None, after=None, container=None,
          underline=False, italic=False, before=None):
        p = (container or doc).add_paragraph()
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if text:
            r = p.add_run(text); r.bold = bold; r.underline = underline; r.italic = italic
            if has_ar(text):
                r.font.name = AR_DOCX_FONT
                r._element.rPr.rFonts.set(qn("w:cs"), AR_DOCX_FONT)
            if size:
                r.font.size = Pt(size)
        return p

    def KVTABLE(rows, container=None, wlab=None):
        wlab = wlab or min(max(len(a) for a, _ in rows) * 0.22 + 0.6, 6.5)
        t = (container or doc).add_table(rows=0, cols=3); t.autofit = False
        for lab, val in rows:
            cells = t.add_row().cells
            cells[0].width, cells[1].width = Cm(wlab), Cm(0.5)
            cells[2].width = Cm(max(W - wlab - 0.5, 3))
            for cell, txt in ((cells[0], lab), (cells[1], ":"), (cells[2], val)):
                pp = cell.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
                pp.add_run(txt)
        return t

    # ---------- kop (teks selalu di tengah; logo di kiri tanpa menggeser teks) ----------
    if k and (k["name"] or k["logo"]):
        box = None
        if k["logo"]:
            lw = KOP_LOGO_W
            t = doc.add_table(rows=1, cols=3); t.autofit = False
            c0, c1, c2 = t.rows[0].cells
            c0.width, c1.width, c2.width = Cm(lw), Cm(W - 2 * lw), Cm(lw)
            c0.paragraphs[0].add_run().add_picture(io.BytesIO(k["logo"]), width=Cm(2.3))
            for cell in (c0, c1, c2):
                _valign_center(cell)
            box = c1
        first = box.paragraphs[0] if box is not None else doc.add_paragraph()
        first.alignment = WA.CENTER
        first.paragraph_format.space_after = Pt(0)
        r = first.add_run(k["name"].upper()); r.bold = True; r.font.size = Pt(c["size"] + 4)
        for line in (k["addr"], k["contact"]):
            if line:
                P(line, WA.CENTER, size=c["size"] - 1, after=0, container=box)
        _hr(P(after=10), double=True)

    # ---------- header memo ----------
    if d.get("memo"):
        m = d["memo"]
        P(d["judul"].upper(), WA.CENTER, bold=True, size=c["size"] + 3, after=0)
        if d.get("nomor_judul"):
            P("No. " + d["nomor_judul"], WA.CENTER, size=c["size"] - 1, after=10)
        rows = [(lab, val) for lab, val in
                (("Tanggal", m.get("tanggal")), ("Dari", m.get("dari")),
                 ("Kepada", m.get("kepada")), ("Tembusan", m.get("tembusan")),
                 ("Subyek", m.get("subyek"))) if val]
        KVTABLE(rows, wlab=2.5)
        P(after=6)
    # ---------- blok nomor / judul ----------
    elif d["meta"]:
        m = d["meta"]
        if m.get("kota_tgl"):
            P(m["kota_tgl"], WA.RIGHT, after=6)
        rows = [(lab, val) for lab, val in
                (("Nomor", m.get("nomor")), ("Sifat", m.get("sifat")),
                 ("Lampiran", m.get("lampiran")), ("Hal", m.get("hal"))) if val]
        if rows:
            KVTABLE(rows, wlab=2.2)
            P(after=4)
        if m.get("tujuan"):
            P("Kepada Yth.", after=0)
            if m.get("up"):
                P("Up. " + m["up"], after=0)
            for ln in tujuan_lines(m):
                P(ln, after=0)
            P("di Tempat")
    else:
        p = P(d["judul"].upper(), WA.CENTER, bold=True, size=c["size"] + 2,
              underline=not d.get("nomor_judul"), after=2 if d.get("nomor_judul") else 12)
        if d.get("nomor_judul"):
            P("Nomor: " + d["nomor_judul"], WA.CENTER, bold=True, after=12)
        if d.get("tentang"):
            P("Tentang", WA.CENTER, after=0)
            P(d["tentang"].upper(), WA.CENTER, bold=True, after=12)

    # ---------- badan ----------
    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                p = P(txt, bold=True, after=4, before=10)
                p.paragraph_format.keep_with_next = True
            elif kind == "li":
                p = doc.add_paragraph(txt, style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
            elif kind == "no":
                lab, body = txt
                t = doc.add_table(rows=1, cols=2); t.autofit = False
                cells = t.rows[0].cells
                cells[0].width, cells[1].width = Cm(0.9), Cm(W - 0.9)
                for cell, s_ in ((cells[0], lab), (cells[1], body)):
                    pp = cell.paragraphs[0]; pp.paragraph_format.space_after = Pt(2)
                    pp.alignment = body_al if cell is cells[1] else WA.LEFT
                    pp.add_run(s_)
            elif kind == "kv":
                KVTABLE(txt); P(after=4)
            elif kind == "tbl":
                rows = txt
                t = doc.add_table(rows=0, cols=len(rows[0]))
                t.alignment = WD_TABLE_ALIGNMENT.LEFT
                for i, row in enumerate(rows):
                    cells = t.add_row().cells
                    for j, val in enumerate(row[:len(cells)]):
                        pp = cells[j].paragraphs[0]
                        pp.paragraph_format.space_after = Pt(0)
                        r = pp.add_run(val); r.bold = (i == 0)
                        if i == 0:
                            pp.alignment = WA.CENTER
                _borders(t)
                P(after=6)
            else:
                P(txt, body_al)

    # ---------- tanda tangan ----------
    ttds = [t for t in d.get("ttds", []) if t.get("nama") or t.get("jabatan")]
    if ttds:
        P(after=0)
        if d.get("verifikasi"):
            P(d["verifikasi"] if d["verifikasi"].lower().startswith(("mengetahui", "menyetujui"))
              else "Mengetahui,", WA.CENTER, after=0)
            P(d["verifikasi"], WA.CENTER, after=0)
            for _ in range(3):
                P(after=0)
            P("(............................................)", WA.CENTER, after=10)

        kota_line = d.get("kota_tgl", "")
        if d.get("tutup_tgl") == "ditetapkan" and kota_line:
            kota, _, tg = kota_line.partition(", ")
            KVTABLE([("Ditetapkan di", kota), ("Pada tanggal", tg)], wlab=3.2)
            P(after=6)
            kota_line = ""

        ncol = min(len(ttds), 2) if len(ttds) != 3 else 3
        rows = [ttds[i:i + ncol] for i in range(0, len(ttds), ncol)]
        for row in rows:
            heads = _heads(row, ncol, kota_line, len(ttds), d.get("materai"))
            t = doc.add_table(rows=1, cols=ncol); t.autofit = False
            for idx in range(ncol):
                cell = t.rows[0].cells[idx]
                cell.width = Cm(W / ncol)
                sg = row[idx] if idx < len(row) else None
                pp = cell.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
                if not sg:
                    continue
                pp.alignment = WA.CENTER
                lines = heads[idx]
                first = True
                for ln in lines:
                    if first:
                        pp.add_run(ln); first = False
                    else:
                        P(ln, WA.CENTER, after=0, container=cell)
                if sg.get("img"):
                    p = P(align=WA.CENTER, after=0, container=cell)
                    p.add_run().add_picture(io.BytesIO(sg["img"]), width=Cm(3.2))
                else:
                    for _ in range(3):
                        P(after=0, container=cell)
                P(sg.get("nama", ""), WA.CENTER, bold=True, underline=True, after=0,
                  container=cell)
                if sg.get("nip"):
                    P(sg["nip"], WA.CENTER, after=0, container=cell)
            P(after=8)

    # ---------- saksi ----------
    if d.get("saksi"):
        P("Saksi-saksi:", bold=True, before=8, after=4)
        for i, s_ in enumerate(d["saksi"], 1):
            P(f"{i}. {s_}", after=0)
            P("   ..............................", after=6)

    # ---------- registrasi ----------
    if d.get("registrasi"):
        t = doc.add_table(rows=1, cols=2); t.autofit = False
        for idx, lines in enumerate((("Dicatat Nomor : ...........", "Tanggal : ..........."),
                                     ("Dicatat di Buku Register", "Nomor : ...........",
                                      "Tanggal : ..........."))):
            cell = t.rows[0].cells[idx]; cell.width = Cm(W / 2)
            for j, ln in enumerate(lines):
                pp = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                pp.paragraph_format.space_after = Pt(0); pp.add_run(ln)
        P(after=8)

    # ---------- tembusan & NB ----------
    if d.get("tembusan"):
        P("Tembusan:", before=10, after=2)
        for i, x in enumerate(d["tembusan"], 1):
            P(f"{i}. {x}", after=0)
    if d.get("nb"):
        P("NB: " + d["nb"], italic=True, before=10)

    # ---------- slip sobek ----------
    if d.get("slip"):
        p = P(before=14, after=6)
        pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr")
        ln = OxmlElement("w:bottom")
        for kk, vv in (("val", "dashed"), ("sz", "6"), ("space", "1"), ("color", "808080")):
            ln.set(qn("w:" + kk), vv)
        bd.append(ln); pPr.append(bd)
        for kind, txt in blocks(d["slip"]):
            if kind == "h":
                P(txt.upper(), WA.CENTER, bold=True, after=6)
            elif kind == "kv":
                KVTABLE(txt)
            elif kind == "li":
                doc.add_paragraph(txt, style="List Bullet")
            else:
                P(txt, body_al)

    # ---------- nomor halaman ----------
    tpl = PAGE_FMT[c["pagefmt"]]
    if tpl:
        p = s.footer.paragraphs[0]
        p.alignment = {"Kiri": WA.LEFT, "Tengah": WA.CENTER, "Kanan": WA.RIGHT}[c["pagepos"]]
        for part in re.split(r"(\{p\}|\{n\})", tpl):
            if part == "{p}":
                _field(p.add_run(), "PAGE")
            elif part == "{n}":
                _field(p.add_run(), "NUMPAGES")
            elif part:
                p.add_run(part)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


# ===================================================================== PDF
def _img(data, w_cm):
    iw, ih = PILImage.open(io.BytesIO(data)).size
    return Image(io.BytesIO(data), width=w_cm * cm, height=w_cm * cm * ih / iw)


def _canvas(tpl, pos, font, size, pw, mx, my):
    class NC(rl_canvas.Canvas):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw); self._st = []

        def showPage(self):
            self._st.append(dict(self.__dict__)); self._startPage()

        def save(self):
            total = len(self._st)
            for x in self._st:
                self.__dict__.update(x)
                if tpl:
                    self.setFont(font, size - 1)
                    t, y = tpl.format(p=self._pageNumber, n=total), my / 2
                    if pos == "Kiri":
                        self.drawString(mx, y, t)
                    elif pos == "Kanan":
                        self.drawRightString(pw - mx, y, t)
                    else:
                        self.drawCentredString(pw / 2, y, t)
                super().showPage()
            super().save()
    return NC


NOPAD = [("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
         ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
         ("VALIGN", (0, 0), (-1, -1), "TOP")]


def build_pdf(d):
    c, k = d["cfg"], d["kop"]
    reg, bold = FONTS[c["font"]]
    pw, ph = [x * cm for x in PAPER[c["paper"]]]
    mx, my = c["mx"] * cm, c["my"] * cm
    W = pw - 2 * mx
    lead = c["size"] * c["spacing"] * 1.2

    def S(name, **kw):
        return ParagraphStyle(name, fontName=kw.pop("fontName", reg),
                              fontSize=kw.pop("fontSize", c["size"]),
                              leading=kw.pop("leading", lead),
                              spaceAfter=kw.pop("spaceAfter", 6), **kw)

    body = S("body", alignment=ALIGN[c["align"]][0])
    ctr = S("ctr", alignment=TA_CENTER, spaceAfter=0)
    small = S("sm", fontSize=c["size"] - 3, alignment=TA_CENTER, spaceAfter=0)
    head = S("hd", fontName=bold, spaceBefore=10, spaceAfter=4)
    story = []

    def KVT(rows, wlab=None):
        wlab = (wlab or min(max(len(a) for a, _ in rows) * 0.22 + 0.6, 6.5)) * cm
        t = Table([[Paragraph(esc(a), body), Paragraph(":", body), Paragraph(esc(b), body)]
                   for a, b in rows],
                  colWidths=[wlab, 0.5 * cm, max(W - wlab - 0.5 * cm, 3 * cm)], hAlign="LEFT")
        t.setStyle(TableStyle(NOPAD))
        return t

    # ---------- kop (teks selalu di tengah; logo di kiri tanpa menggeser teks) ----------
    if k and (k["name"] or k["logo"]):
        txt = [Paragraph(esc(k["name"].upper()),
                         S("kn", fontName=bold, fontSize=c["size"] + 4, alignment=TA_CENTER,
                           leading=(c["size"] + 4) * 1.2, spaceAfter=0))]
        for line in (k["addr"], k["contact"]):
            if line:
                txt.append(Paragraph(esc(line), S("kl", fontSize=c["size"] - 1,
                                                  alignment=TA_CENTER, spaceAfter=0)))
        if k["logo"]:
            lw = KOP_LOGO_W * cm
            logo = _img(k["logo"], 2.3)
            logo.hAlign = "LEFT"
            t = Table([[logo, txt, ""]], colWidths=[lw, W - 2 * lw, lw])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
            story.append(t)
        else:
            story += txt
        story += [Spacer(1, 4), HRFlowable(width="100%", thickness=1.6, color=colors.black),
                  Spacer(1, 1.5), HRFlowable(width="100%", thickness=0.6, color=colors.black),
                  Spacer(1, 12)]

    if d.get("memo"):
        m = d["memo"]
        story.append(Paragraph(esc(d["judul"].upper()),
                               S("t", fontName=bold, fontSize=c["size"] + 3,
                                 alignment=TA_CENTER, spaceAfter=2)))
        if d.get("nomor_judul"):
            story.append(Paragraph("No. " + esc(d["nomor_judul"]), small))
        rows = [(a, b) for a, b in (("Tanggal", m.get("tanggal")), ("Dari", m.get("dari")),
                                    ("Kepada", m.get("kepada")),
                                    ("Tembusan", m.get("tembusan")),
                                    ("Subyek", m.get("subyek"))) if b]
        story += [Spacer(1, 10), KVT(rows, 2.5), Spacer(1, 10)]
    elif d["meta"]:
        m = d["meta"]
        if m.get("kota_tgl"):
            story.append(Paragraph(esc(m["kota_tgl"]), S("r", alignment=2)))
        rows = [(a, b) for a, b in (("Nomor", m.get("nomor")), ("Sifat", m.get("sifat")),
                                    ("Lampiran", m.get("lampiran")), ("Hal", m.get("hal")))
                if b]
        if rows:
            story += [KVT(rows, 2.2), Spacer(1, 8)]
        if m.get("tujuan"):
            lines = ["Kepada Yth."]
            if m.get("up"):
                lines.append("Up. " + m["up"])
            lines += tujuan_lines(m) + ["di Tempat"]
            for ln in lines[:-1]:
                story.append(Paragraph(esc(ln), S("y", spaceAfter=0)))
            story.append(Paragraph(esc(lines[-1]), body))
    else:
        story.append(Paragraph(f"<u>{esc(d['judul'].upper())}</u>" if not d.get("nomor_judul")
                               else esc(d["judul"].upper()),
                               S("t", fontName=bold, fontSize=c["size"] + 2,
                                 alignment=TA_CENTER, leading=(c["size"] + 2) * 1.3,
                                 spaceAfter=2 if d.get("nomor_judul") else 14)))
        if d.get("nomor_judul"):
            story.append(Paragraph("Nomor: " + esc(d["nomor_judul"]),
                                   S("tn", fontName=bold, alignment=TA_CENTER, spaceAfter=14)))
        if d.get("tentang"):
            story.append(Paragraph("Tentang", ctr))
            story.append(Paragraph(esc(d["tentang"].upper()),
                                   S("tt", fontName=bold, alignment=TA_CENTER, spaceAfter=14)))

    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                story.append(Paragraph(esc(txt), head))
            elif kind == "li":
                story.append(Paragraph(esc(txt), S("li", alignment=ALIGN[c["align"]][0],
                                                   spaceAfter=3, leftIndent=14),
                                       bulletText="\u2022"))
            elif kind == "no":
                lab, bodytxt = txt
                t = Table([[Paragraph(esc(lab), body), Paragraph(esc(bodytxt), body)]],
                          colWidths=[0.9 * cm, W - 0.9 * cm], hAlign="LEFT")
                t.setStyle(TableStyle(NOPAD + [("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                story.append(t)
            elif kind == "kv":
                story += [KVT(txt), Spacer(1, 5)]
            elif kind == "tbl":
                rows = txt
                ncol = max(len(r) for r in rows)
                cw = W / ncol
                data = [[Paragraph(f"<b>{esc(x)}</b>" if i == 0 else esc(x),
                                   S("c", alignment=TA_CENTER if i == 0 else TA_LEFT,
                                     spaceAfter=0))
                         for x in (r + [""] * (ncol - len(r)))] for i, r in enumerate(rows)]
                t = Table(data, colWidths=[cw] * ncol, hAlign="LEFT", repeatRows=1)
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.Color(.90, .90, .90)),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 4),
                                       ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                                       ("TOPPADDING", (0, 0), (-1, -1), 3),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                story += [t, Spacer(1, 8)]
            else:
                story.append(Paragraph(esc(txt), body))

    ttds = [t for t in d.get("ttds", []) if t.get("nama") or t.get("jabatan")]
    if ttds:
        story.append(Spacer(1, 10))
        if d.get("verifikasi"):
            story += [Paragraph("Mengetahui,", ctr), Paragraph(esc(d["verifikasi"]), ctr),
                      Spacer(1, 1.6 * cm),
                      Paragraph("(............................................)", ctr),
                      Spacer(1, 12)]
        kota_line = d.get("kota_tgl", "")
        if d.get("tutup_tgl") == "ditetapkan" and kota_line:
            kota, _, tg = kota_line.partition(", ")
            story += [KVT([("Ditetapkan di", kota), ("Pada tanggal", tg)], 3.2), Spacer(1, 8)]
            kota_line = ""

        ncol = min(len(ttds), 2) if len(ttds) != 3 else 3
        for i in range(0, len(ttds), ncol):
            row = ttds[i:i + ncol]
            heads = _heads(row, ncol, kota_line, len(ttds), d.get("materai"))
            cells = []
            for idx in range(ncol):
                if idx >= len(row):
                    cells.append([Spacer(1, 1)]); continue
                sg = row[idx]
                cell = [Paragraph(esc(ln) or "&nbsp;", ctr) for ln in heads[idx]]
                cell.append(_img(sg["img"], 3.2) if sg.get("img") else Spacer(1, 1.8 * cm))
                cell.append(Paragraph(f"<u><b>{esc(sg.get('nama', ''))}</b></u>", ctr))
                if sg.get("nip"):
                    cell.append(Paragraph(esc(sg["nip"]), ctr))
                cells.append(cell)
            t = Table([cells], colWidths=[W / ncol] * ncol, hAlign="LEFT")
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 2),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
            story += [KeepTogether(t), Spacer(1, 10)]

    if d.get("saksi"):
        story.append(Paragraph("<b>Saksi-saksi:</b>", S("sk", spaceBefore=8)))
        for i, s_ in enumerate(d["saksi"], 1):
            story.append(Paragraph(f"{i}. {esc(s_)}", S("s1", spaceAfter=0)))
            story.append(Paragraph("&nbsp;&nbsp;&nbsp;..............................", body))

    if d.get("registrasi"):
        left = [Paragraph("Dicatat Nomor : ...........", ctr),
                Paragraph("Tanggal : ...........", ctr)]
        right = [Paragraph("Dicatat di Buku Register", ctr),
                 Paragraph("Nomor : ...........", ctr), Paragraph("Tanggal : ...........", ctr)]
        t = Table([[left, right]], colWidths=[W / 2] * 2, hAlign="LEFT")
        t.setStyle(TableStyle(NOPAD))
        story += [Spacer(1, 8), t]

    if d.get("tembusan"):
        story.append(Paragraph("<b>Tembusan:</b>", S("tb", spaceBefore=12, spaceAfter=2)))
        for i, x in enumerate(d["tembusan"], 1):
            story.append(Paragraph(f"{i}. {esc(x)}", S("tb2", spaceAfter=0)))
    if d.get("nb"):
        story.append(Paragraph(f"<i>NB: {esc(d['nb'])}</i>", S("nb", spaceBefore=12)))

    if d.get("slip"):
        story += [Spacer(1, 16),
                  HRFlowable(width="100%", thickness=0.8, color=colors.grey, dash=(3, 3)),
                  Spacer(1, 10)]
        for kind, txt in blocks(d["slip"]):
            if kind == "h":
                story.append(Paragraph(esc(txt.upper()),
                                       S("sh", fontName=bold, alignment=TA_CENTER,
                                         spaceAfter=8)))
            elif kind == "kv":
                story.append(KVT(txt))
            elif kind == "li":
                story.append(Paragraph(esc(txt), S("sl", spaceAfter=2), bulletText="\u2022"))
            else:
                story.append(Paragraph(esc(txt), body))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=(pw, ph), leftMargin=mx, rightMargin=mx,
                      topMargin=my, bottomMargin=my).build(
        story, canvasmaker=_canvas(PAGE_FMT[c["pagefmt"]], c["pagepos"], reg, c["size"],
                                   pw, mx, my))
    return buf.getvalue()


# ============================================================ PRATINJAU HTML
def _b64(data):
    return "data:image/png;base64," + base64.b64encode(data).decode()


def build_html(d):
    c, k = d["cfg"], d["kop"]
    fam = {"Times New Roman": "'Times New Roman', Times, serif",
           "Arial": "Arial, Helvetica, sans-serif",
           "Calibri": "Calibri, Carlito, sans-serif"}[c["font"]]
    pw = PAPER[c["paper"]][0]
    just = "justify" if c["align"].startswith("Rata kiri-kanan") else "left"
    o = []

    def kvhtml(rows, cls="kv"):
        r = "".join(f'<tr><td class="kvl">{esc(a)}</td><td class="kvc">:</td>'
                    f'<td>{esc(b)}</td></tr>' for a, b in rows)
        return f'<table class="{cls}">{r}</table>'

    if k and (k["name"] or k["logo"]):
        sub = "".join(f'<div class="kl">{esc(x)}</div>' for x in (k["addr"], k["contact"]) if x)
        logo = (f'<div class="klogo"><img src="{_b64(k["logo"])}" style="width:2.3cm"></div>'
                if k["logo"] else "")
        cls = "kop haslogo" if k["logo"] else "kop"
        o.append(f'<div class="{cls}">{logo}'
                 f'<div class="ktxt"><div class="kn">{esc(k["name"].upper())}</div>'
                 f'{sub}</div></div><hr>')

    if d.get("memo"):
        m = d["memo"]
        o.append(f'<div class="judul nomargin">{esc(d["judul"].upper())}</div>')
        if d.get("nomor_judul"):
            o.append(f'<div class="nomor">No. {esc(d["nomor_judul"])}</div>')
        o.append(kvhtml([(a, b) for a, b in
                         (("Tanggal", m.get("tanggal")), ("Dari", m.get("dari")),
                          ("Kepada", m.get("kepada")), ("Tembusan", m.get("tembusan")),
                          ("Subyek", m.get("subyek"))) if b]))
    elif d["meta"]:
        m = d["meta"]
        if m.get("kota_tgl"):
            o.append(f'<div class="right">{esc(m["kota_tgl"])}</div>')
        rows = [(a, b) for a, b in (("Nomor", m.get("nomor")), ("Sifat", m.get("sifat")),
                                    ("Lampiran", m.get("lampiran")), ("Hal", m.get("hal")))
                if b]
        if rows:
            o.append(kvhtml(rows))
        if m.get("tujuan"):
            tj = ["Kepada Yth."] + (["Up. " + m["up"]] if m.get("up") else []) \
                 + tujuan_lines(m) + ["di Tempat"]
            o.append('<div class="tuj">' + "<br>".join(esc(x) for x in tj) + "</div>")
    else:
        u = "" if d.get("nomor_judul") else "underline"
        o.append(f'<div class="judul {u}">{esc(d["judul"].upper())}</div>')
        if d.get("nomor_judul"):
            o.append(f'<div class="nomor"><b>Nomor: {esc(d["nomor_judul"])}</b></div>')
        if d.get("tentang"):
            o.append(f'<div class="ctr">Tentang</div>'
                     f'<div class="judul">{esc(d["tentang"].upper())}</div>')

    for part in (d["pembuka"], d["isi"], d["penutup"]):
        for kind, txt in blocks(part):
            if kind == "h":
                o.append(f'<div class="h">{esc(txt)}</div>')
            elif kind == "li":
                o.append(f'<ul><li>{esc(txt)}</li></ul>')
            elif kind == "no":
                o.append(f'<table class="no"><tr><td class="nol">{esc(txt[0])}</td>'
                         f'<td>{esc(txt[1])}</td></tr></table>')
            elif kind == "kv":
                o.append(kvhtml(txt))
            elif kind == "tbl":
                rows = txt
                ncol = max(len(r) for r in rows)
                body_rows = ""
                for i, r in enumerate(rows):
                    tag = "th" if i == 0 else "td"
                    cells = "".join(f"<{tag}>{esc(x)}</{tag}>"
                                    for x in (r + [""] * (ncol - len(r))))
                    body_rows += f"<tr>{cells}</tr>"
                o.append(f'<table class="grid">{body_rows}</table>')
            else:
                o.append(f"<p>{esc(txt)}</p>")

    ttds = [t for t in d.get("ttds", []) if t.get("nama") or t.get("jabatan")]
    if ttds:
        if d.get("verifikasi"):
            o.append(f'<div class="sgwrap" style="justify-content:center"><div class="sg">'
                     f'<div>Mengetahui,</div><div>{esc(d["verifikasi"])}</div>'
                     f'<div style="height:1.6cm"></div>'
                     f'<div>(............................................)</div></div></div>')
        kota_line = d.get("kota_tgl", "")
        if d.get("tutup_tgl") == "ditetapkan" and kota_line:
            kota, _, tg = kota_line.partition(", ")
            o.append(kvhtml([("Ditetapkan di", kota), ("Pada tanggal", tg)]))
            kota_line = ""
        ncol = min(len(ttds), 2) if len(ttds) != 3 else 3
        for i in range(0, len(ttds), ncol):
            row = ttds[i:i + ncol]
            heads = _heads(row, ncol, kota_line, len(ttds), d.get("materai"))
            cells = ""
            for idx in range(ncol):
                if idx >= len(row):
                    cells += "<td></td>"; continue
                sg = row[idx]
                inner = "".join(f"<div>{esc(ln) or '&nbsp;'}</div>" for ln in heads[idx])
                inner += (f'<img src="{_b64(sg["img"])}" style="width:3.2cm">'
                          if sg.get("img") else '<div style="height:1.8cm"></div>')
                inner += f'<div class="nm">{esc(sg.get("nama", ""))}</div>'
                if sg.get("nip"):
                    inner += f"<div>{esc(sg['nip'])}</div>"
                cells += f'<td class="sgcell">{inner}</td>'
            o.append(f'<table class="sgrow"><tr>{cells}</tr></table>')

    if d.get("saksi"):
        o.append("<div class='h'>Saksi-saksi:</div>")
        for i, s_ in enumerate(d["saksi"], 1):
            o.append(f"<div>{i}. {esc(s_)}</div>"
                     "<div style='margin:0 0 8pt 1em'>..............................</div>")
    if d.get("registrasi"):
        o.append('<table class="sgrow"><tr>'
                 '<td class="sgcell">Dicatat Nomor : ...........<br>Tanggal : ...........</td>'
                 '<td class="sgcell">Dicatat di Buku Register<br>Nomor : ...........<br>'
                 'Tanggal : ...........</td></tr></table>')
    if d.get("tembusan"):
        o.append('<div class="h">Tembusan:</div>')
        o.append("".join(f"<div>{i}. {esc(x)}</div>" for i, x in enumerate(d["tembusan"], 1)))
    if d.get("nb"):
        o.append(f'<p><i>NB: {esc(d["nb"])}</i></p>')
    if d.get("slip"):
        o.append('<div class="cut"></div>')
        for kind, txt in blocks(d["slip"]):
            if kind == "h":
                o.append(f'<div class="judul">{esc(txt.upper())}</div>')
            elif kind == "kv":
                o.append(kvhtml(txt))
            elif kind == "li":
                o.append(f"<ul><li>{esc(txt)}</li></ul>")
            else:
                o.append(f"<p>{esc(txt)}</p>")

    return f"""<style>
.paper{{background:#fff;color:#111;font-family:{fam};font-size:{c["size"]}pt;
line-height:{c["spacing"]};padding:{c["my"]}cm {c["mx"]}cm;width:{pw}cm;max-width:100%;
box-sizing:border-box;margin:0 auto;box-shadow:0 2px 14px rgba(0,0,0,.25);border-radius:4px;}}
.paper p{{text-align:{just};margin:0 0 6pt;}}
.paper hr{{border:0;border-top:1.6px solid #000;border-bottom:.6px solid #000;height:2px;
margin:5px 0 12px;}}
.paper .kop{{position:relative;text-align:center;}}
.paper .kop.haslogo{{display:flex;align-items:center;justify-content:center;
padding:0 {KOP_LOGO_W}cm;min-height:2.4cm;box-sizing:border-box;}}
.paper .klogo{{position:absolute;left:0;top:50%;transform:translateY(-50%);}}
.paper .ktxt{{width:100%;text-align:center;}}
.paper .kn{{font-weight:700;font-size:{c["size"] + 4}pt;}}
.paper .kl{{font-size:{c["size"] - 1}pt;}}
.paper .right{{text-align:right;margin-bottom:6pt;}}
.paper .ctr{{text-align:center;}}
.paper .tuj{{margin:10pt 0;}}
.paper .judul{{text-align:center;font-weight:700;font-size:{c["size"] + 2}pt;margin:0 0 12pt;}}
.paper .judul.underline{{text-decoration:underline;}}
.paper .judul.nomargin{{margin-bottom:2pt;}}
.paper .nomor{{text-align:center;margin-bottom:12pt;}}
.paper .h{{font-weight:700;margin:10pt 0 4pt;}}
.paper ul{{margin:0 0 4pt 0;padding-left:1.2em;}}
.paper table.kv,.paper table.no{{margin:0 0 5pt;border-collapse:collapse;}}
.paper table.kv td,.paper table.no td{{padding:0 0 1pt;vertical-align:top;}}
.paper table.kv .kvl{{white-space:nowrap;padding-right:.3cm;}}
.paper table.kv .kvc{{width:.4cm;}}
.paper table.no .nol{{width:.9cm;vertical-align:top;}}
.paper table.grid{{border-collapse:collapse;width:100%;margin:0 0 8pt;}}
.paper table.grid th,.paper table.grid td{{border:.6px solid #000;padding:3pt 4pt;
vertical-align:top;font-size:{c["size"] - 1}pt;}}
.paper table.grid th{{background:#e6e6e6;text-align:center;}}
.paper table.sgrow{{width:100%;margin-top:12pt;border-collapse:collapse;}}
.paper .sgcell{{text-align:center;vertical-align:top;padding:0 4pt;}}
.paper .sgcell .nm{{font-weight:700;text-decoration:underline;}}
.paper .materai{{border:1px solid #888;width:2.2cm;margin:4pt auto;font-size:
{max(c["size"] - 4, 6)}pt;color:#666;}}
.paper .sgwrap{{display:flex;margin-top:12pt;}}
.paper .sg{{width:7cm;text-align:center;}}
.paper .cut{{border-top:1px dashed #999;margin:16pt 0 10pt;}}
</style><div class="paper">{"".join(o)}</div>"""
